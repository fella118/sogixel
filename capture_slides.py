#!/usr/bin/env python3
"""
Capture all 8 slides from the Sogixel presentation and export as PDF + PPTX.
Uses Playwright for high-quality browser screenshots.
"""

import subprocess
import sys
import os
import time

# Install playwright if needed
try:
    from playwright.sync_api import sync_playwright
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--break-system-packages", "playwright"])
    subprocess.check_call([sys.executable, "-m", "playwright", "install", "chromium"])
    from playwright.sync_api import sync_playwright

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from PIL import Image
import img2pdf

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SCREENSHOTS_DIR = os.path.join(BASE_DIR, "screenshots")
os.makedirs(SCREENSHOTS_DIR, exist_ok=True)

TOTAL_SLIDES = 8
WIDTH = 1920
HEIGHT = 1080

def capture_slides():
    """Capture each slide as a high-res PNG screenshot."""
    print("🎬 Capturing slides...")
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page(viewport={"width": WIDTH, "height": HEIGHT})
        page.goto("http://localhost:8888/index.html")
        page.wait_for_load_state("networkidle")
        time.sleep(1)  # let animations finish

        for i in range(TOTAL_SLIDES):
            slide_num = i + 1
            path = os.path.join(SCREENSHOTS_DIR, f"slide_{slide_num:02d}.png")
            # Wait for current slide animation to settle
            time.sleep(0.8)
            page.screenshot(path=path, type="png")
            print(f"  ✅ Slide {slide_num} captured -> {path}")
            if i < TOTAL_SLIDES - 1:
                page.keyboard.press("ArrowRight")
                time.sleep(1)  # wait for transition

        browser.close()
    print("📸 All slides captured!\n")


def create_pdf():
    """Create PDF from slide screenshots."""
    print("📄 Creating PDF...")
    pdf_path = os.path.join(BASE_DIR, "Sogixel_Presentation.pdf")
    
    image_paths = [
        os.path.join(SCREENSHOTS_DIR, f"slide_{i:02d}.png")
        for i in range(1, TOTAL_SLIDES + 1)
    ]
    
    # Convert PNGs to a single PDF
    with open(pdf_path, "wb") as f:
        f.write(img2pdf.convert(image_paths, layout_fun=img2pdf.get_layout_fun(
            pagesize=(img2pdf.mm_to_pt(338.667), img2pdf.mm_to_pt(190.5))  # 16:9 ratio
        )))
    
    print(f"  ✅ PDF saved -> {pdf_path}\n")


def create_pptx():
    """Create PPTX from slide screenshots."""
    print("📊 Creating PPTX...")
    pptx_path = os.path.join(BASE_DIR, "Sogixel_Presentation.pptx")
    
    prs = Presentation()
    # Set widescreen 16:9 slide size
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_layout = prs.slide_layouts[6]  # blank layout
    
    for i in range(1, TOTAL_SLIDES + 1):
        img_path = os.path.join(SCREENSHOTS_DIR, f"slide_{i:02d}.png")
        slide = prs.slides.add_slide(blank_layout)
        
        # Add image as full-slide background
        slide.shapes.add_picture(
            img_path,
            left=Emu(0),
            top=Emu(0),
            width=prs.slide_width,
            height=prs.slide_height
        )
    
    prs.save(pptx_path)
    print(f"  ✅ PPTX saved -> {pptx_path}\n")


if __name__ == "__main__":
    capture_slides()
    create_pdf()
    create_pptx()
    print("🎉 Done! Both files saved in:", BASE_DIR)
